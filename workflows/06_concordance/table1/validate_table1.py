#!/usr/bin/env python3
"""Validate the generated Table 1 rows and optional Word-table cells.

This publication validator was created during the v1.1.0 recovery pass. It does
not modify the manuscript. Pass ``--manuscript`` only when the protected Word
file is locally available and python-docx is installed.
"""

from __future__ import annotations

import argparse
import csv
import math
import tempfile
from pathlib import Path

from generate_table1 import main as generate_table


ROOT = Path(__file__).resolve().parents[3]
JSON_GCSV = ROOT / "figures/figure_03_imputation/impute_gcsv.json"
JSON_RSQ = ROOT / "figures/figure_03_imputation/rsquare_all16.json"
SUMMARY = ROOT / "data/summary/imputation/imputation_summary.csv"
PROVENANCE = ROOT / "docs/validation/table1_cell_provenance.tsv"

NUMERIC_FIELDS = ("dosage_r2_maf01", "bestguess_r2_maf01", "NRD_pct")


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def normalize_strategy(value: str) -> str:
    return "Uncorrected" if value == "Raw (uncorrected)" else value


def validate_generated_rows() -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    with tempfile.TemporaryDirectory(prefix="ort-table1-") as temp_dir:
        generated_path = Path(temp_dir) / "table1.csv"
        generate_table(str(JSON_GCSV), str(JSON_RSQ), str(generated_path))
        generated = read_csv(generated_path)
    expected = read_csv(SUMMARY)
    if len(generated) != 16 or len(expected) != 16:
        raise AssertionError(
            f"Table 1 row count changed: generated={len(generated)}, expected={len(expected)}"
        )
    for row_number, (observed, reference) in enumerate(
        zip(generated, expected), start=1
    ):
        if observed["sample"] != reference["sample"]:
            raise AssertionError(f"sample differs in Table 1 row {row_number}")
        if normalize_strategy(observed["strategy"]) != reference["strategy"]:
            raise AssertionError(f"strategy differs in Table 1 row {row_number}")
        for field in NUMERIC_FIELDS:
            if not math.isclose(
                float(observed[field]),
                float(reference[field]),
                rel_tol=0.0,
                abs_tol=5e-5,
            ):
                raise AssertionError(
                    f"{field} differs in Table 1 row {row_number}: "
                    f"{observed[field]} != {reference[field]}"
                )
    return generated, expected


def validate_published_provenance() -> None:
    with PROVENANCE.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle, delimiter="\t"))
    if len(rows) != 80:
        raise AssertionError(f"expected 80 Table 1 provenance cells, found {len(rows)}")
    for row in rows:
        if row["status"] != "PASS":
            raise AssertionError("published Table 1 cell provenance contains a failure")
        if row["source_file"] != "data/summary/imputation/imputation_summary.csv":
            raise AssertionError("published Table 1 provenance source path changed")


def validate_word_table(manuscript: Path, expected: list[dict[str, str]]) -> None:
    try:
        from docx import Document
    except ImportError as exc:
        raise SystemExit("python-docx is required with --manuscript") from exc

    document = Document(manuscript)
    if len(document.tables) != 1:
        raise AssertionError(f"expected one manuscript table, found {len(document.tables)}")
    table = document.tables[0]
    if (len(table.rows), len(table.columns)) != (17, 5):
        raise AssertionError(
            f"expected a 17x5 manuscript table, found {len(table.rows)}x{len(table.columns)}"
        )
    fields = ("sample", "strategy", *NUMERIC_FIELDS)
    matched = 0
    for row_number, (word_row, reference) in enumerate(
        zip(table.rows[1:], expected), start=1
    ):
        values = [cell.text.strip() for cell in word_row.cells]
        for field, observed in zip(fields, values):
            if field in NUMERIC_FIELDS:
                passed = math.isclose(
                    float(observed),
                    float(reference[field]),
                    rel_tol=0.0,
                    abs_tol=5e-5,
                )
            else:
                passed = observed == reference[field]
            if not passed:
                raise AssertionError(
                    f"Word Table 1 row {row_number}, {field}: "
                    f"{observed!r} != {reference[field]!r}"
                )
            matched += 1
    if matched != 80:
        raise AssertionError(f"expected 80 matched Word cells, found {matched}")
    print("PASS  Table 1 Word cells: 80/80")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--manuscript",
        type=Path,
        help="Optional protected Reduce16 Word manuscript for the 80-cell check",
    )
    args = parser.parse_args()
    _, expected = validate_generated_rows()
    validate_published_provenance()
    print("PASS  Table 1 generated rows: 16/16")
    print("PASS  Published Table 1 cell provenance: 80/80")
    if args.manuscript is not None:
        validate_word_table(args.manuscript, expected)


if __name__ == "__main__":
    main()
